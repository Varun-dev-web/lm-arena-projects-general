"""Supporting documents: executive summary, financial summary, agreement, Q&A, script."""
import os
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from notes import NOTES

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "PROJECT_FUTURE_Supporting_Documents.docx")

NAVY = RGBColor(0x0A, 0x13, 0x22)
BLUE = RGBColor(0x1D, 0x4E, 0xD8)
SLATE = RGBColor(0x4E, 0x5F, 0x78)
MUTED = RGBColor(0x8A, 0x99, 0xAE)
CYAN = RGBColor(0x0E, 0x9C, 0xB8)


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcPr.append(shd)


def set_cell_margins(table, top=70, bottom=70, left=110, right=110):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)


def hairline_borders(table, color="E2E8F1"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    for edge in ("insideH",):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblPr.append(borders)


def para(doc, text="", size=10.5, color=NAVY, bold=False, space_before=0, space_after=6,
         align="l", line=1.35, italic=False, font="Inter", track=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = line
    pf.alignment = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER,
                    "r": WD_ALIGN_PARAGRAPH.RIGHT, "j": WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    if text:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        if track:
            rpr = r._element.get_or_add_rPr()
            rpr.set(qn("w:spacing"), str(int(track * 20)))
    return p


def rich(p, chunks):
    """chunks: list of (text, dict(size,color,bold,italic))"""
    for text, fmt in chunks:
        r = p.add_run(text)
        r.font.size = Pt(fmt.get("size", 10.5))
        r.font.color.rgb = fmt.get("color", NAVY)
        r.font.bold = fmt.get("bold", False)
        r.font.italic = fmt.get("italic", False)
        r.font.name = "Inter"
    return p


def eyebrow(doc, text, size=8, color=CYAN, space_before=16, space_after=3):
    return para(doc, text.upper(), size=size, color=color, bold=True, space_before=space_before,
                space_after=space_after, track=1.6)


def h1(doc, text, space_before=6):
    p = para(doc, text, size=19, color=NAVY, bold=True, space_before=space_before,
             space_after=4, line=1.12)
    return p


def h2(doc, text, space_before=12):
    return para(doc, text, size=11.5, color=BLUE, bold=True, space_before=space_before,
                space_after=3, line=1.2)


def rule(doc, color="E2E8F1", space_before=2, space_after=10):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)
    return p


def bullet(doc, text, size=10.5, color=NAVY, space_after=3, glyph="—"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Cm(0.6)
    pf.first_line_indent = Cm(-0.6)
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.3
    r = p.add_run(f"{glyph}   ")
    r.font.size = Pt(size)
    r.font.color.rgb = MUTED
    r.font.name = "Inter"
    r2 = p.add_run(text)
    r2.font.size = Pt(size)
    r2.font.color.rgb = color
    r2.font.name = "Inter"
    return p


def table(doc, rows, widths, header=True, size=9.5, header_bg="0A1322", header_color="FFFFFF",
          zebra=None, align_last_right=True, bold_last_row=False):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    hairline_borders(t)
    set_cell_margins(t)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = t.cell(r_i, c_i)
            cell.width = Cm(widths[c_i])
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.2
            is_head = header and r_i == 0
            is_total = bold_last_row and r_i == len(rows) - 1
            if c_i > 0 and align_last_right:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            run = p.add_run(str(val))
            run.font.size = Pt(size)
            run.font.name = "Inter"
            run.font.bold = is_head or is_total
            run.font.color.rgb = RGBColor.from_string(header_color) if is_head else NAVY
            if is_head:
                set_cell_bg(cell, header_bg)
            elif is_total:
                set_cell_bg(cell, "F2F6FF")
            elif zebra and r_i % 2 == 0:
                set_cell_bg(cell, "F7F9FC")
    for c_i, w in enumerate(widths):
        for r in t.rows:
            r.cells[c_i].width = Cm(w)
    return t


def add_page_number_footer(section, label):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(f"{label}   ·   ")
    run.font.size = Pt(7.5)
    run.font.color.rgb = MUTED
    run.font.name = "Inter"
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), "15")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "8A99AE")
    rPr.append(sz); rPr.append(color)
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


# ---------------------------------------------------------------------------
# document
# ---------------------------------------------------------------------------
def build():
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Inter"
    style.font.size = Pt(10.5)
    style.font.color.rgb = NAVY
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Inter")

    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    add_page_number_footer(sec, "PROJECT FUTURE  ·  SUPPORTING DOCUMENTS")

    # ---------------------------------------------------------------- cover
    para(doc, "PROJECT FUTURE", size=8.5, color=CYAN, bold=True, track=2.2, space_before=10,
         space_after=2)
    para(doc, "Personal Development & Technology Fund", size=26, color=NAVY, bold=True,
         space_after=2, line=1.05)
    para(doc, "Supporting documents for a three-month pilot proposal", size=12, color=SLATE,
         space_after=10)
    para(doc, "Prepared for my parents  ·  October 2026", size=9.5, color=MUTED, space_after=6)
    rule(doc, space_after=12)

    para(doc, "This document accompanies the presentation. It contains a one-page summary, the full "
              "financial breakdown, the accountability agreement, a longer list of questions and "
              "answers, and the complete speaking script.",
         size=10, color=SLATE, space_after=10)

    table(doc, [
        ["Contents", ""],
        ["1", "Executive summary"],
        ["2", "Financial summary"],
        ["3", "Accountability agreement"],
        ["4", "Questions parents may ask (14)"],
        ["5", "Full presentation script"],
    ], widths=[2.0, 14.0], size=9.5, zebra=True)

    # ------------------------------------------------------- 1. exec summary
    doc.add_page_break()
    eyebrow(doc, "Section 1", space_before=0)
    h1(doc, "Executive summary")
    para(doc, "The idea in one page.", size=10.5, color=SLATE, space_after=10)

    h2(doc, "The proposal", space_before=4)
    para(doc, "A recurring monthly development fund of ₹2,000, provided at the beginning of each "
              "month. ₹1,900 is allocated to a Future Development Fund that accumulates over time; "
              "₹100 is allocated to ordinary personal expenses.",
         size=10.5, space_after=8)

    h2(doc, "Why")
    para(doc, "Robotics, electronics and programming are continuous rather than occasional: they "
              "need parts, tools and repeated experiments. Equipment is currently bought through "
              "individual requests. A monthly fund would let those requirements be planned, "
              "prioritised and accumulated instead of decided one purchase at a time. It "
              "supplements existing parental support; it does not replace it.",
         size=10.5, space_after=8)

    h2(doc, "The structure")
    bullet(doc, "₹2,000 per month, paid at the start of the month.")
    bullet(doc, "₹1,900 to the Future Development Fund — accumulated, not automatically spent.")
    bullet(doc, "₹100 for personal expenses.")
    bullet(doc, "A shared ledger recording every contribution, purchase and balance.")
    bullet(doc, "An accountability agreement covering academics, school, home and personal discipline.")
    bullet(doc, "A formal joint review at the end of month three.")

    h2(doc, "The ask")
    para(doc, "Approval for a three-month pilot: ₹6,000 in total, of which ₹5,700 would sit in the "
              "development fund and ₹300 would be personal allocation. At the end of the pilot both "
              "sides review the arrangement and decide whether to continue, modify or discontinue it.",
         size=10.5, space_after=8)

    h2(doc, "What is deliberately not claimed")
    para(doc, "No interest, investment return or financial appreciation is included in any figure. "
              "No equipment prices, academic results or project outcomes are assumed. The "
              "projections describe contributions only, and phase timing depends on readiness rather "
              "than dates.",
         size=10.5, space_after=8)

    # ---------------------------------------------------- 2. financial summary
    doc.add_page_break()
    eyebrow(doc, "Section 2", space_before=0)
    h1(doc, "Financial summary")
    para(doc, "Every figure below is a contribution, not an investment return.", size=10.5,
         color=SLATE, space_after=10)

    h2(doc, "Monthly allocation", space_before=4)
    table(doc, [
        ["Allocation", "Amount", "Share"],
        ["Future Development Fund", "₹1,900", "95%"],
        ["Personal expenses", "₹100", "5%"],
        ["Total per month", "₹2,000", "100%"],
    ], widths=[8.6, 4.0, 3.4], bold_last_row=True, zebra=True)
    para(doc, "The development allocation is not spent by default. Its primary purpose is gradual "
              "accumulation towards educational resources, robotics components, tools and, in time, "
              "larger equipment.", size=9.5, color=SLATE, space_before=6, space_after=8)

    h2(doc, "Projected accumulation (no withdrawals)")
    table(doc, [
        ["Period", "Contributed", "Development fund", "Personal allocation"],
        ["3 months", "₹6,000", "₹5,700", "₹300"],
        ["6 months", "₹12,000", "₹11,400", "₹600"],
        ["9 months", "₹18,000", "₹17,100", "₹900"],
        ["12 months", "₹24,000", "₹22,800", "₹1,200"],
    ], widths=[4.0, 4.0, 4.4, 3.6], zebra=True)
    para(doc, "Calculated as ₹1,900 × months and ₹100 × months. No interest, returns or appreciation "
              "are included; the balance is only what has been contributed.", size=9.5, color=SLATE,
         space_before=6, space_after=8)

    h2(doc, "The three-month pilot")
    table(doc, [
        ["Item", "Amount"],
        ["Parental contribution (3 × ₹2,000)", "₹6,000"],
        ["Future Development Fund (3 × ₹1,900)", "₹5,700"],
        ["Personal allocation (3 × ₹100)", "₹300"],
    ], widths=[11.0, 5.0], zebra=True)

    h2(doc, "How the fund would be recorded")
    para(doc, "A single ledger, visible to both parents at any time, recording: contributions "
              "received; current balance; each purchase and its purpose; remaining balance; and "
              "progress towards the equipment goals. Significant purchases would be discussed "
              "before any money is spent.",
         size=10, space_after=6)

    # -------------------------------------------------- 3. accountability
    doc.add_page_break()
    eyebrow(doc, "Section 3", space_before=0)
    h1(doc, "Accountability agreement")
    para(doc, "A three-month arrangement, reviewed together.", size=10.5, color=SLATE,
         space_after=10)

    para(doc, "In return for a monthly development fund of ₹2,000, I agree to the commitments below "
              "for the duration of the three-month pilot. The commitments are assessed on "
              "consistency, effort and progress rather than on a single result.",
         size=10.5, space_after=10)

    areas = [
        ("1. Academic performance", [
            "Maintain consistent academic performance across subjects.",
            "Complete assignments and academic responsibilities on time.",
            "Show continuous improvement over the term.",
            "Keep a balanced approach to learning rather than last-minute effort.",
        ]),
        ("2. School responsibilities", [
            "Fulfil responsibilities as Emerald Vice Captain.",
            "Participate responsibly in school activities and events.",
            "Demonstrate leadership and reliability in the house.",
        ]),
        ("3. Household responsibilities", [
            "Contribute to agreed household activities.",
            "Complete agreed responsibilities without needing reminders.",
            "Demonstrate cooperation and responsibility at home.",
        ]),
        ("4. Personal discipline", [
            "Maintain effective time management around deadlines.",
            "Balance academics and technical interests.",
            "Use technology productively.",
            "Maintain consistency in personal goals and routines.",
        ]),
    ]
    for title, items in areas:
        h2(doc, title, space_before=10)
        for it in items:
            bullet(doc, it, size=9.8)

    h2(doc, "Review arrangement", space_before=12)
    para(doc, "A short monthly check-in, and a formal review at the end of month three covering:",
         size=10, space_after=4)
    for i, c in enumerate(["Academic consistency",
                           "Fulfilment of school and household responsibilities",
                           "Personal discipline",
                           "Financial transparency",
                           "Progress in technical learning",
                           "Whether the arrangement is working effectively"], 1):
        bullet(doc, f"{i}. {c}", size=9.8, glyph="")

    para(doc, "Following the review, the decision rests with my parents: continue as it stands, "
              "continue with changes, or discontinue the arrangement.",
         size=10, space_before=6, space_after=14)

    # signature block
    t = doc.add_table(rows=2, cols=2)
    no_borders(t)
    labels = [("Student", "Date"), ("Parent / Guardian", "Date")]
    for i, (left, right) in enumerate(labels):
        for j, label in enumerate((left, right)):
            cell = t.cell(i, j)
            cell.width = Cm(8.6)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run("_______________________________")
            r.font.size = Pt(10)
            r.font.color.rgb = MUTED
            r.font.name = "Inter"
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_after = Pt(6)
            r2 = p2.add_run(label)
            r2.font.size = Pt(8)
            r2.font.color.rgb = SLATE
            r2.font.name = "Inter"

    # ------------------------------------------------------------ 4. Q&A
    doc.add_page_break()
    eyebrow(doc, "Section 4", space_before=0)
    h1(doc, "Questions parents may ask")
    para(doc, "Fourteen questions, answered plainly. The first five also appear in the presentation.",
         size=10.5, color=SLATE, space_after=10)

    qa = [
        ("1. Why do you need money every month when we already buy the equipment you ask for?",
         "The aim isn't to replace what you already do. It's to add structure: a fixed amount, "
         "agreed in advance, that lets me plan purchases and accumulate towards something larger "
         "instead of deciding each one in the moment."),
        ("2. Why can't you simply ask us whenever you need something?",
         "Individual requests work for immediate needs. They don't work for something that costs "
         "more than a month's allowance, because there's nothing accumulating in the background. "
         "A fund solves the second problem without removing the first."),
        ("3. How can we be certain the money will be used properly?",
         "One ledger records every contribution, purchase, purpose and balance, and you have access "
         "to it at any time. Anything significant is discussed before it's bought, and the "
         "three-month review exists precisely so you can judge the record rather than my intentions."),
        ("4. What happens if academic performance or discipline declines?",
         "It becomes a review item. The agreement treats academics as consistency, effort and "
         "progress rather than a single exam mark, and school, household responsibilities and "
         "discipline are reviewed at the same time. If something slips, we would discuss it in the "
         "monthly check-in rather than waiting three months."),
        ("5. Do you genuinely need expensive equipment right now?",
         "No, and I'm not asking for it. The point of the pilot is to plan and save now so that a "
         "larger purchase later is deliberate and researched rather than sudden."),
        ("6. What exactly will the ₹100 be spent on?",
         "Ordinary personal expenses — nothing to do with equipment. Keeping it separate means the "
         "development fund stays clean and easy to check."),
        ("7. What if you don't spend the ₹1,900 in a month?",
         "Then nothing happens: it stays in the fund. The intention is that most months it simply "
         "accumulates, and that the fund is spent deliberately rather than monthly."),
        ("8. Isn't this just an allowance with a different name?",
         "An allowance is for spending. This is mostly for saving — 95% of it is expected to sit "
         "there and grow towards equipment, with a written record and a review attached."),
        ("9. What if we want to give a smaller amount?",
         "That's a reasonable outcome. ₹1,500 a month, for example, would still build roughly "
         "₹4,275 in the development fund over three months. A version you're comfortable with is "
         "more useful than the version I designed."),
        ("10. Who holds the money?",
         "Whichever arrangement you prefer — a separate account, a dedicated wallet, or an amount "
         "you hold and release against the ledger. The record and the review matter more than the "
         "container."),
        ("11. What happens at the end of three months?",
         "We review six things together: academics, responsibilities, discipline, transparency, "
         "technical progress, and whether the arrangement is working. Then you decide whether to "
         "continue, change it, or stop."),
        ("12. What if you lose interest in robotics?",
         "Then the fund stops being useful and you should say so. The review exists for that. I'd "
         "rather the arrangement be judged honestly every three months than be assumed permanent."),
        ("13. How much time will this take away from studies?",
         "The time is already being spent on projects; the fund changes how they're paid for, not "
         "how much time they take. The accountability agreement is where time balance is actually "
         "reviewed."),
        ("14. What if we say no?",
         "Then I'll keep doing what I'm doing now, and I'll ask the way I always have. I'd still "
         "like to keep the written record, because it's useful either way."),
    ]
    for q, a in qa:
        para(doc, q, size=10.2, bold=True, space_before=8, space_after=2, line=1.25)
        para(doc, a, size=10, color=SLATE, space_after=4, line=1.35)

    # ---------------------------------------------------------- 5. script
    doc.add_page_break()
    eyebrow(doc, "Section 5", space_before=0)
    h1(doc, "Full presentation script")
    para(doc, "Approximately 6 minutes when spoken at a natural pace. Slide numbers match the "
              "presentation file.", size=10.5, color=SLATE, space_after=10)

    titles = {
        1: "Title — Project Future",
        2: "The Vision",
        3: "Why a structured fund",
        4: "The Proposal",
        5: "The Financial Roadmap",
        6: "The Equipment Goals",
        7: "The Development Journey",
        8: "The Accountability Agreement",
        9: "Complete Transparency",
        10: "The Three-Month Pilot",
        11: "Addressing Parental Concerns",
        12: "The Proposal — Decision Requested",
        13: "Closing",
    }
    for i in range(1, 14):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r1 = p.add_run(f"SLIDE {i:02d}   ")
        r1.font.size = Pt(8)
        r1.font.bold = True
        r1.font.color.rgb = BLUE
        r1.font.name = "Inter"
        r2 = p.add_run(titles[i])
        r2.font.size = Pt(8)
        r2.font.color.rgb = MUTED
        r2.font.name = "Inter"
        r2.font.bold = True
        for chunk in NOTES[i].split("\n\n"):
            para(doc, chunk.strip(), size=10.2, color=NAVY, space_after=5, line=1.42)

    para(doc, "— End —", size=9, color=MUTED, align="c", space_before=18)

    doc.core_properties.title = "PROJECT FUTURE — Supporting Documents"
    doc.core_properties.subject = "Executive summary, financial summary, accountability agreement, Q&A and script"
    doc.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    build()
