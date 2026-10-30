"""
PROJECT FUTURE — deck engine.

A small layout DSL that renders every primitive twice:
  1. into a python-pptx presentation (the deliverable)
  2. into a PIL image (pixel-accurate preview, used for visual QC)

Both backends use the same Inter font metrics, so what the preview
shows is what PowerPoint shows.
"""
from __future__ import annotations

import copy
import os
from lxml import etree
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ----------------------------------------------------------------------------
# palette
# ----------------------------------------------------------------------------
NAVY_950 = "#050A14"
NAVY_900 = "#070D1B"
NAVY_850 = "#0A1223"
NAVY_800 = "#0D1830"
NAVY_700 = "#152341"
NAVY_600 = "#1E3054"
INK = "#0A1322"
SLATE = "#4E5F78"
MUTED = "#8A99AE"
LINE = "#E2E8F1"
LINE_DARK = "#1B2A46"
BG = "#FFFFFF"
BG_SOFT = "#F7F9FC"

BLUE = "#1D4ED8"        # electric blue (deep)
BLUE_BRIGHT = "#2563EB"
BLUE_LIGHT = "#5B8DEF"
CYAN = "#12B8D6"
CYAN_BRIGHT = "#22D3EE"
WHITE = "#FFFFFF"

FONT_NAME = "Inter"
FONTS = {
    300: "/home/user/.fonts/inter/Inter-300.ttf",
    400: "/home/user/.fonts/inter/Inter-400.ttf",
    500: "/home/user/.fonts/inter/Inter-500.ttf",
    600: "/home/user/.fonts/inter/Inter-600.ttf",
    700: "/home/user/.fonts/inter/Inter-700.ttf",
    800: "/home/user/.fonts/inter/Inter-800.ttf",
}
# PowerPoint can only address regular / bold through a typeface name, so the
# deck uses exactly two weights and the preview mirrors that mapping.
WEIGHT_MAP = {300: 400, 400: 400, 500: 400, 600: 700, 700: 700, 800: 700}

PX = 192.0          # preview pixels per inch
LINE_FACTOR = 1.21  # Inter hhea (ascent + descent) / em


def rgb(hexstr: str) -> tuple[int, int, int]:
    hexstr = hexstr.lstrip("#")
    return tuple(int(hexstr[i : i + 2], 16) for i in (0, 2, 4))


def rgba(hexstr: str, alpha: float) -> tuple[int, int, int, int]:
    return rgb(hexstr) + (int(round(max(0.0, min(1.0, alpha)) * 255)),)


_font_cache: dict[tuple[int, int], ImageFont.FreeTypeFont] = {}


def pil_font(size_pt: float, weight: int) -> ImageFont.FreeTypeFont:
    w = WEIGHT_MAP.get(weight, 400)
    px = int(round(size_pt * PX / 72.0))
    key = (px, w)
    if key not in _font_cache:
        _font_cache[key] = ImageFont.truetype(FONTS[w], px)
    return _font_cache[key]


def text_width(s: str, size_pt: float, weight: int, track_pt: float = 0.0) -> float:
    """Width in inches."""
    f = pil_font(size_pt, weight)
    px = f.getlength(s) + (track_pt * PX / 72.0) * max(0, len(s) - 1)
    return px / PX


def wrap_text(s: str, size_pt: float, weight: int, width_in: float, track_pt: float = 0.0) -> list[str]:
    out: list[str] = []
    for raw in s.split("\n"):
        words = raw.split(" ")
        cur = ""
        for word in words:
            trial = word if not cur else cur + " " + word
            if text_width(trial, size_pt, weight, track_pt) <= width_in or not cur:
                cur = trial
            else:
                out.append(cur)
                cur = word
        out.append(cur)
    return out


# ----------------------------------------------------------------------------
# geometry helpers
# ----------------------------------------------------------------------------
def _emo(x):
    return Emu(int(round(x * 914400)))


def _set_alpha(color_obj, alpha: float):
    """Attach <a:alpha> to a python-pptx colour element."""
    srgb = color_obj._xFill.find(qn("a:srgbClr")) if hasattr(color_obj, "_xFill") else None
    if srgb is None:
        return
    for old in srgb.findall(qn("a:alpha")):
        srgb.remove(old)
    el = etree.SubElement(srgb, qn("a:alpha"))
    el.set("val", str(int(round(alpha * 100000))))


def _set_gradfill(shape, stops, angle_deg=90.0, radial=False):
    """stops: [(pos 0-1, hex, alpha)]"""
    spPr = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        for el in spPr.findall(qn(tag)):
            spPr.remove(el)
    xml_stops = []
    for pos, hexc, alpha in stops:
        a = "" if alpha >= 0.999 else f'<a:alpha val="{int(round(alpha * 100000))}"/>'
        xml_stops.append(f'<a:gs pos="{int(round(pos * 100000))}"><a:srgbClr val="{hexc.lstrip("#")}">{a}</a:srgbClr></a:gs>')
    if radial:
        direction = '<a:path path="circle"><a:fillToRect l="50000" t="50000" r="50000" b="50000"/></a:path>'
    else:
        # `angle_deg` uses the preview convention (0 = gradient runs upward, 90 = to the
        # right). DrawingML measures clockwise from +x, hence the -90 offset.
        ang = int(round(((angle_deg - 90.0) % 360) * 60000))
        direction = f'<a:lin ang="{ang}" scaled="0"/>'
    xml = (
        '<a:gradFill xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'rotWithShape="1"><a:gsLst>' + "".join(xml_stops) + "</a:gsLst>" + direction + "</a:gradFill>"
    )
    frag = etree.fromstring(xml)
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(frag)
    else:
        spPr.append(frag)


def _set_shadow(shape, blur=0.16, dist=0.05, alpha=0.10, color=INK, direction=90):
    spPr = shape._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    xml = (
        f'<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:outerShdw blurRad="{int(blur * 914400)}" dist="{int(dist * 914400)}" '
        f'dir="{int(direction * 60000)}" algn="ctr" rotWithShape="0">'
        f'<a:srgbClr val="{color.lstrip("#")}"><a:alpha val="{int(alpha * 100000)}"/></a:srgbClr>'
        f"</a:outerShdw></a:effectLst>"
    )
    spPr.append(etree.fromstring(xml))


DASH_PPTX = {
    "dash": MSO_LINE_DASH_STYLE.DASH,
    "dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
    "sysdash": MSO_LINE_DASH_STYLE.SQUARE_DOT,
    "lgdash": MSO_LINE_DASH_STYLE.LONG_DASH,
}


class Slide:
    def __init__(self, deck: "Deck", bg: str, dark: bool):
        self.deck = deck
        self.dark = dark
        self.bg = bg
        self.prs_slide = deck.prs.slides.add_slide(deck.blank)
        self.img = Image.new("RGBA", (int(13.333 * PX), int(7.5 * PX)), rgba(bg, 1.0))
        self.draw = ImageDraw.Draw(self.img, "RGBA")
        self.warnings: list[str] = []
        self._cover_background(bg)

    # -- background ---------------------------------------------------------
    def _cover_background(self, bg):
        fill = self.prs_slide.background.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor.from_string(bg.lstrip("#"))
        # PIL image starts opaque because it was created with RGBA(hex,255)

    # -- low level ----------------------------------------------------------
    def _shape(self, kind, x, y, w, h, fill=None, line=None, lw=1.0, dash=None,
               alpha=1.0, line_alpha=1.0, r=None, rot=0.0, shadow=None):
        shp = self.prs_slide.shapes.add_shape(kind, _emo(x), _emo(y), _emo(w), _emo(h))
        shp.shadow.inherit = False
        if r is not None:
            try:
                shp.adjustments[0] = max(0.0, min(0.5, r / max(0.001, min(w, h))))
            except Exception:
                pass
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
            if alpha < 0.999:
                _set_alpha(shp.fill.fore_color, alpha)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = RGBColor.from_string(line.lstrip("#"))
            shp.line.width = Pt(lw)
            if line_alpha < 0.999:
                _set_alpha(shp.line.color, line_alpha)
            if dash:
                shp.line.dash_style = DASH_PPTX[dash]
        if rot:
            shp.rotation = rot
        if shadow:
            _set_shadow(shp, **shadow)
        return shp

    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, r=0.0, dash=None,
             alpha=1.0, line_alpha=1.0, shadow=None, rot=0.0):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if r and r > 0 else MSO_SHAPE.RECTANGLE
        self._shape(kind, x, y, w, h, fill, line, lw, dash, alpha, line_alpha,
                    r if r else None, rot, shadow)
        # preview
        box = [x * PX, y * PX, (x + w) * PX, (y + h) * PX]
        rad = (r * PX) if r else 0
        if shadow:
            self._preview_shadow(box, rad, shadow)
        if alpha >= 0.999 and line_alpha >= 0.999:
            d, lay = self.draw, None
        else:
            lay, d = self._layer()
        if fill is not None:
            d.rounded_rectangle(box, radius=rad, fill=rgba(fill, alpha))
        if line is not None:
            d.rounded_rectangle(box, radius=rad, outline=rgba(line, line_alpha),
                                width=max(1, int(round(lw * PX / 72))))
        if lay is not None:
            self.img.alpha_composite(lay)
        return self

    def _layer(self):
        img = Image.new("RGBA", self.img.size, (0, 0, 0, 0))
        return img, ImageDraw.Draw(img, "RGBA")

    def _preview_shadow(self, box, rad, shadow):
        layer = Image.new("RGBA", self.img.size, (0, 0, 0, 0))
        d = ImageDraw.Draw(layer, "RGBA")
        off_x = 0
        off_y = shadow.get("dist", 0.05) * PX
        b = [box[0] + off_x, box[1] + off_y, box[2] + off_x, box[3] + off_y]
        d.rounded_rectangle(b, radius=rad, fill=rgba(shadow.get("color", INK), shadow.get("alpha", 0.1)))
        layer = layer.filter(ImageFilter.GaussianBlur(shadow.get("blur", 0.16) * PX / 1.9))
        self.img.alpha_composite(layer)

    def line(self, x1, y1, x2, y2, color, lw=1.0, dash=None, alpha=1.0, cap="flat"):
        conn = self.prs_slide.shapes.add_connector(1, _emo(x1), _emo(y1), _emo(x2), _emo(y2))
        conn.line.color.rgb = RGBColor.from_string(color.lstrip("#"))
        conn.line.width = Pt(lw)
        if alpha < 0.999:
            _set_alpha(conn.line.color, alpha)
        if dash:
            conn.line.dash_style = DASH_PPTX[dash]
        wpx = max(1, int(round(lw * PX / 72)))
        if alpha >= 0.999:
            d, lay = self.draw, None
        else:
            lay, d = self._layer()
        if dash:
            self._preview_dashed((x1 * PX, y1 * PX), (x2 * PX, y2 * PX), color, alpha, wpx, dash, d)
        else:
            d.line([x1 * PX, y1 * PX, x2 * PX, y2 * PX], fill=rgba(color, alpha), width=wpx)
            if cap == "round":
                for (px, py) in ((x1, y1), (x2, y2)):
                    d.ellipse([px * PX - wpx / 2, py * PX - wpx / 2,
                               px * PX + wpx / 2, py * PX + wpx / 2],
                              fill=rgba(color, alpha))
        if lay is not None:
            self.img.alpha_composite(lay)
        return self

    def _preview_dashed(self, p0, p1, color, alpha, wpx, style, d=None):
        d = d or self.draw
        import math
        dash_len = {"dash": 8, "dot": 3, "sysdash": 6, "lgdash": 14}[style] * max(1, wpx)
        gap = dash_len * 0.85
        x0, y0 = p0
        x1, y1 = p1
        total = math.hypot(x1 - x0, y1 - y0)
        if total <= 0:
            return
        ux, uy = (x1 - x0) / total, (y1 - y0) / total
        pos = 0.0
        while pos < total:
            end = min(pos + dash_len, total)
            d.line([x0 + ux * pos, y0 + uy * pos, x0 + ux * end, y0 + uy * end],
                   fill=rgba(color, alpha), width=wpx)
            pos = end + gap

    def circle(self, cx, cy, r, fill=None, line=None, lw=1.0, alpha=1.0, line_alpha=1.0,
               shadow=None, radial=None):
        if radial:
            shp = self._shape(MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r)
            _set_gradfill(shp, radial, radial=True)
            self._preview_radial(cx, cy, r, radial)
            return self
        shp = self._shape(MSO_SHAPE.OVAL, cx - r, cy - r, 2 * r, 2 * r, fill, line, lw,
                          None, alpha, line_alpha, None, 0.0, shadow)
        box = [(cx - r) * PX, (cy - r) * PX, (cx + r) * PX, (cy + r) * PX]
        if shadow:
            self._preview_shadow(box, (box[2] - box[0]) / 2, shadow)
        if alpha >= 0.999 and line_alpha >= 0.999:
            d, lay = self.draw, None
        else:
            lay, d = self._layer()
        if fill is not None:
            d.ellipse(box, fill=rgba(fill, alpha))
        if line is not None:
            d.ellipse(box, outline=rgba(line, line_alpha), width=max(1, int(round(lw * PX / 72))))
        if lay is not None:
            self.img.alpha_composite(lay)
        return self

    def _preview_radial(self, cx, cy, r, stops):
        import numpy as np
        size = int(2 * r * PX) + 2
        ys, xs = np.mgrid[0:size, 0:size]
        dist = np.sqrt((xs - size / 2) ** 2 + (ys - size / 2) ** 2) / (size / 2)
        dist = np.clip(dist, 0, 1)
        arr = np.zeros((size, size, 4), dtype=np.uint8)
        positions = [s[0] for s in stops]
        for i in range(size):
            pass
        # piecewise linear interpolation across stops
        cols = np.array([rgb(s[1]) for s in stops], dtype=float)
        alps = np.array([s[2] for s in stops], dtype=float)
        pos = np.array(positions, dtype=float)
        r_ch = np.interp(dist, pos, cols[:, 0])
        g_ch = np.interp(dist, pos, cols[:, 1])
        b_ch = np.interp(dist, pos, cols[:, 2])
        a_ch = np.interp(dist, pos, alps) * 255
        arr[..., 0] = r_ch
        arr[..., 1] = g_ch
        arr[..., 2] = b_ch
        arr[..., 3] = a_ch
        layer = Image.fromarray(arr, "RGBA")
        self.img.alpha_composite(layer, (int((cx - r) * PX), int((cy - r) * PX)))

    def grad_rect(self, x, y, w, h, stops, angle=90.0, radial=False, shape_radius=0.0):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if shape_radius else MSO_SHAPE.RECTANGLE
        shp = self.prs_slide.shapes.add_shape(kind, _emo(x), _emo(y), _emo(w), _emo(h))
        shp.shadow.inherit = False
        shp.line.fill.background()
        if shape_radius:
            shp.adjustments[0] = shape_radius / max(0.001, min(w, h))
        _set_gradfill(shp, stops, angle, radial)
        if radial:
            self._preview_radial(x + w / 2, y + h / 2, max(w, h) / 2, stops)
        else:
            self._preview_linear(x, y, w, h, stops, angle)
        return self

    def _preview_linear(self, x, y, w, h, stops, angle):
        import numpy as np
        # 256px ramp along the gradient axis, then rotate
        ramp = np.zeros((1, 256, 4), dtype=np.uint8)
        positions = np.array([s[0] for s in stops])
        cols = np.array([rgb(s[1]) for s in stops], dtype=float)
        alps = np.array([s[2] for s in stops], dtype=float)
        t = np.linspace(0, 1, 256)
        for ch in range(3):
            ramp[0, :, ch] = np.interp(t, positions, cols[:, ch]).astype(np.uint8)
        ramp[0, :, 3] = (np.interp(t, positions, alps) * 255).astype(np.uint8)
        grad = Image.fromarray(ramp, "RGBA").resize((256, 256), Image.BILINEAR)
        import math
        ang = math.radians(angle)  # clockwise from top
        size = int(max(w, h) * PX * 2)
        big = grad.resize((size, size), Image.BILINEAR)
        big = big.rotate(-(angle - 90), resample=Image.BICUBIC, expand=False)
        px_w, px_h = int(w * PX), int(h * PX)
        left = int((size - px_w) / 2)
        top = int((size - px_h) / 2)
        crop = big.crop((left, top, left + px_w, top + px_h))
        self.img.alpha_composite(crop, (int(x * PX), int(y * PX)))

    def poly(self, pts, fill=None, line=None, lw=1.0, alpha=1.0, line_alpha=1.0, close=True):
        builder = self.prs_slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=914400)
        builder.add_line_segments(pts[1:], close=close)
        shp = builder.convert_to_shape()
        shp.shadow.inherit = False
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
            if alpha < 0.999:
                _set_alpha(shp.fill.fore_color, alpha)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = RGBColor.from_string(line.lstrip("#"))
            shp.line.width = Pt(lw)
            if line_alpha < 0.999:
                _set_alpha(shp.line.color, line_alpha)
        else:
            shp.line.fill.background()
        px = [(p[0] * PX, p[1] * PX) for p in pts]
        if alpha >= 0.999 and line_alpha >= 0.999:
            d, lay = self.draw, None
        else:
            lay, d = self._layer()
        if fill:
            d.polygon(px, fill=rgba(fill, alpha))
        if line:
            seq = px + [px[0]] if close else px
            d.line(seq, fill=rgba(line, line_alpha), width=max(1, int(round(lw * PX / 72))),
                   joint="curve")
        if lay is not None:
            self.img.alpha_composite(lay)
        return self

    def pie(self, cx, cy, r, start_deg, end_deg, fill, alpha=1.0):
        shp = self._shape(MSO_SHAPE.PIE, cx - r, cy - r, 2 * r, 2 * r, fill=fill, alpha=alpha)
        shp.adjustments[0] = start_deg
        shp.adjustments[1] = end_deg
        box = [(cx - r) * PX, (cy - r) * PX, (cx + r) * PX, (cy + r) * PX]
        self.draw.pieslice(box, start=start_deg, end=end_deg, fill=rgba(fill, alpha))
        return self

    # -- text ---------------------------------------------------------------
    def text(self, x, y, w, h, s, size=12, weight=400, color=INK, align="l", anchor="t",
             lh=1.3, track=0.0, space_after=0.0, paras=None, caps=False, opacity=1.0,
             name=None, check_overflow=True):
        """Draw text. `s` may be a string (with \n) or a list of paragraph dicts."""
        if isinstance(s, list):
            paras = s
        elif paras is None:
            paras = [{"t": s}]
        else:
            paras = [{"t": s}]

        # normalise paragraphs
        norm = []
        for i, p in enumerate(paras):
            if isinstance(p, str):
                p = {"t": p}
            p = dict(p)
            p.setdefault("size", size)
            p.setdefault("weight", weight)
            p.setdefault("color", color)
            p.setdefault("align", align)
            p.setdefault("lh", lh)
            p.setdefault("track", track)
            p.setdefault("space_before", 0.0)
            p.setdefault("space_after", space_after)
            p.setdefault("caps", caps)
            p.setdefault("opacity", opacity)
            p.setdefault("runs", None)
            norm.append(p)

        # ---- pptx ----
        tb = self.prs_slide.shapes.add_textbox(_emo(x), _emo(y), _emo(w), _emo(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        align_map = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
        first = True
        for p in norm:
            para = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            para.alignment = align_map[p["align"]]
            para.line_spacing = p["lh"]
            if p["space_before"]:
                para.space_before = Pt(p["space_before"])
            if p["space_after"]:
                para.space_after = Pt(p["space_after"])
            runs = p["runs"] or [(p["t"], {})]
            for rtext, rover in runs:
                r = para.add_run()
                r.text = rtext.upper() if (p["caps"] and not rover.get("caps_lock")) else rtext
                f = r.font
                f.name = FONT_NAME
                f.size = Pt(rover.get("size", p["size"]))
                f.bold = WEIGHT_MAP.get(rover.get("weight", p["weight"]), 400) == 700
                f.color.rgb = RGBColor.from_string(rover.get("color", p["color"]).lstrip("#"))
                if p["opacity"] < 0.999:
                    _set_alpha(f.color, p["opacity"])
                tr = rover.get("track", p["track"])
                if tr:
                    r.font._rPr.set("spc", str(int(round(tr * 100))))

        # ---- preview ----
        return self._draw_text_preview(x, y, w, h, norm, anchor, check_overflow, name)

    def _draw_text_preview(self, x, y, w, h, paras, anchor, check_overflow, name):
        layout = []
        total_h = 0.0
        for p in paras:
            size = p["size"]
            lh_px = size * LINE_FACTOR * p["lh"] * PX / 72.0
            track_px = p["track"] * PX / 72.0
            lines = []
            if p.get("runs"):
                # simple run-aware wrap: treat as one string for wrapping
                joined = "".join(r[0] for r in p["runs"])
                lines = wrap_text(joined, size, p["weight"], w, p["track"]) if w > 0 else joined.split("\n")
                # render runs sequentially on a single line assumption
            else:
                txt = p["t"].upper() if p["caps"] else p["t"]
                lines = wrap_text(txt, size, p["weight"], w, p["track"]) if w > 0 else txt.split("\n")
            height = len(lines) * lh_px
            layout.append((p, lines, height, lh_px, track_px))
            total_h += p.get("space_before", 0) * PX / 72.0 + height + p.get("space_after", 0) * PX / 72.0

        cy = {"t": y * PX, "m": y * PX + (h * PX - total_h) / 2, "b": y * PX + h * PX - total_h}[anchor]
        if check_overflow and total_h > h * PX + 0.18 * PX:
            self.warnings.append(
                f"{name or 'text'}: overflow {total_h / PX:.2f}in > {h:.2f}in :: "
                f"{paras[0].get('t','')[:48]!r}"
            )
        for p, lines, height, lh_px, track_px in layout:
            cy += p.get("space_before", 0) * PX / 72.0
            for line in lines:
                f = pil_font(p["size"], p["weight"])
                ascent, descent = f.getmetrics()
                lw = f.getlength(line) + track_px * max(0, len(line) - 1)
                if p["align"] == "l":
                    lx = x * PX
                elif p["align"] == "c":
                    lx = x * PX + (w * PX - lw) / 2
                else:
                    lx = x * PX + w * PX - lw
                baseline = cy + (lh_px - (ascent + descent)) / 2 + ascent
                if p.get("runs"):
                    rx = lx
                    for rtext, rover in p["runs"]:
                        rf = pil_font(rover.get("size", p["size"]), rover.get("weight", p["weight"]))
                        rtrack = rover.get("track", p["track"]) * PX / 72.0
                        self._draw_tracked(rtext, rx, baseline, rf, rover.get("color", p["color"]),
                                           rtrack, p.get("opacity", 1.0))
                        rx += rf.getlength(rtext) + rtrack * max(0, len(rtext) - 1)
                else:
                    self._draw_tracked(line, lx, baseline, f, p["color"], track_px, p.get("opacity", 1.0))
                cy += lh_px
            cy += p.get("space_after", 0) * PX / 72.0
        return total_h / PX

    def _draw_tracked(self, s, lx, baseline, f, color, track_px, opacity=1.0):
        if opacity >= 0.999:
            if track_px <= 0.01:
                self.draw.text((lx, baseline), s, font=f, fill=rgb(color), anchor="ls")
            else:
                cx = lx
                for ch in s:
                    self.draw.text((cx, baseline), ch, font=f, fill=rgb(color), anchor="ls")
                    cx += f.getlength(ch) + track_px
        else:
            layer = Image.new("RGBA", self.img.size, (0, 0, 0, 0))
            d = ImageDraw.Draw(layer, "RGBA")
            d.text((lx, baseline), s, font=f, fill=rgba(color, opacity), anchor="ls")
            self.img.alpha_composite(layer)

    # -- misc ---------------------------------------------------------------
    def notes(self, s: str):
        self.prs_slide.notes_slide.notes_text_frame.text = s
        return self

    def warn(self, msg):
        self.warnings.append(msg)

    def finalize(self):
        return self.img.convert("RGB")


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.blank = self.prs.slide_layouts[6]
        self.slides: list[Slide] = []

    def slide(self, bg=BG, dark=False):
        s = Slide(self, bg, dark)
        self.slides.append(s)
        return s

    def save(self, path):
        self.prs.save(path)

    def report(self):
        for i, s in enumerate(self.slides, 1):
            for w in s.warnings:
                print(f"  [slide {i:02d}] {w}")
